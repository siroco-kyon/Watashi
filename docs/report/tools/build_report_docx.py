from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.part import Part
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17324D"
TEAL = "2A8C82"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F9"
MID_GRAY = "64748B"
BODY = "263238"
WHITE = "FFFFFF"

SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
SCREENSHOT_SPECS = {
    "PANDA-OPERATIONS": {
        "title": "PANDAの現行画面",
        "guide": "代表的な操作手順が分かる画面を挿入\n機密情報・個人情報はマスキングする",
        "caption": "図1 PANDAの現行画面と代表的な操作手順（スクリーンショット差し替え欄）",
        "alt": "PANDAの現行画面を後から挿入するためのプレースホルダー",
    },
    "WATASHI-CLIENT": {
        "title": "Watashi クライアント画面",
        "guide": "二ペイン表示と転送センターが同時に見える画面を挿入\n実データのパス・利用者名は必要に応じてマスキングする",
        "caption": "図6 Watashiの二ペイン表示と転送センター（スクリーンショット差し替え欄）",
        "alt": "Watashiクライアントの二ペイン表示と転送センターのスクリーンショットを後から挿入するためのプレースホルダー",
    },
    "WATASHI-ADMIN": {
        "title": "Watashi ユーザー権限管理画面",
        "guide": "共有・許可サブパス・操作種別が分かる画面を挿入\n本文で説明する設定項目が読める解像度にする",
        "caption": "図7 Watashiのユーザー権限管理画面（スクリーンショット差し替え欄）",
        "alt": "Watashiのユーザー権限管理画面のスクリーンショットを後から挿入するためのプレースホルダー",
    },
    "WATASHI-AUDIT": {
        "title": "Watashi 操作ログ画面",
        "guide": "条件検索・結果一覧・CSV出力が分かる画面を挿入\n利用者名・パス・端末情報は\n合成データ化またはマスキングする",
        "caption": "図5 Watashiの操作ログ検索・結果確認・CSV出力画面（スクリーンショット差し替え欄）",
        "alt": "Watashiの操作ログ検索・結果確認・CSV出力画面のスクリーンショットを後から挿入するためのプレースホルダー",
    },
}


def set_east_asia(run, font_name: str) -> None:
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), font_name)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_toc_with_static_result(doc: Document, entries: list[tuple[str, int]]) -> None:
    """Create an updatable TOC field with a readable cached result."""
    first = doc.add_paragraph()
    begin_run = first.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    begin_run._r.extend([begin, instr, separate])

    for idx, (title, page) in enumerate(entries):
        p = first if idx == 0 else doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.25))
        add_inline(p, title, 10.5)
        r = p.add_run("\t")
        set_east_asia(r, "Yu Gothic")
        r = p.add_run(str(page))
        r.font.size = Pt(10.5)
        set_east_asia(r, "Yu Gothic")

    end_run = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_hyperlink(paragraph, label: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Yu Mincho")
    rpr.extend([rfonts, color, underline])
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


TOKEN_RE = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`)")


def add_inline(paragraph, text: str, font_size: float | None = None) -> None:
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_east_asia(run, "Yu Mincho")
            if font_size:
                run.font.size = Pt(font_size)
        token = match.group(0)
        if token.startswith("[") and "](http" in token:
            label, url = token[1:].split("](", 1)
            add_hyperlink(paragraph, label, url[:-1])
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_east_asia(run, "Yu Gothic")
            if font_size:
                run.font.size = Pt(font_size)
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
            run.font.size = Pt((font_size or 11) - 0.5)
            rpr = run._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF2F6")
            rpr.append(shd)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_east_asia(run, "Yu Mincho")
        if font_size:
            run.font.size = Pt(font_size)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Mincho")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Title", 28, NAVY, 0, 16),
        ("Subtitle", 15, TEAL, 0, 12),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9.5)
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = True

    quote = doc.styles["Quote"]
    quote.font.name = "Calibri"
    quote.font.size = Pt(10.5)
    quote.font.italic = False
    quote.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    quote._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Mincho")
    quote.paragraph_format.left_indent = Inches(0.35)
    quote.paragraph_format.right_indent = Inches(0.15)
    quote.paragraph_format.space_before = Pt(6)
    quote.paragraph_format.space_after = Pt(10)


def add_page_number_footer(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Watashi 技術レポート  |  ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    set_east_asia(run, "Yu Gothic")
    add_field(p, "PAGE", "1")


def add_running_header(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("PANDAの課題を踏まえたWatashiの設計・実装・検証")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    set_east_asia(run, "Yu Gothic")
    p.paragraph_format.space_after = Pt(4)


def configure_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.92)
    section.right_margin = Inches(0.92)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


from report_diagrams import build_diagrams


def build_screenshot_placeholders(asset_dir: Path) -> dict[str, tuple[Path, Path]]:
    from report_diagrams import Figure
    results = {}
    for key, spec in SCREENSHOT_SPECS.items():
        figure = Figure(spec['title'], 1012)
        figure.box(60, 180, 1680, 650, 'スクリーンショット差し替え欄')
        figure.text(900, 470, spec['guide'], 38, center=True)
        results[key] = figure.save(asset_dir, 'placeholder-' + key.lower())
    return results


def attach_svg_fallback(doc: Document, inline_shape, svg_path: Path) -> None:
    package = doc.part.package
    partname = package.next_partname("/word/media/image%d.svg")
    svg_part = Part(partname, "image/svg+xml", svg_path.read_bytes(), package)
    rel_id = doc.part.relate_to(svg_part, RT.IMAGE)
    blips = inline_shape._inline.xpath(".//a:blip")
    if not blips:
        raise RuntimeError(f"Could not locate drawing blip for {svg_path}")
    ext_list = OxmlElement("a:extLst")
    ext = OxmlElement("a:ext")
    ext.set("uri", SVG_EXT_URI)
    svg_blip = etree.Element(f"{{{SVG_NS}}}svgBlip", nsmap={"asvg": SVG_NS})
    svg_blip.set(qn("r:embed"), rel_id)
    ext.append(svg_blip)
    ext_list.append(ext)
    blips[0].append(ext_list)


def add_picture_with_alt(
    doc,
    path: Path,
    width: float,
    alt_text: str,
    svg_path: Path | None = None,
    page_break_before: bool = False,
) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = page_break_before
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(width))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt_text)
    if svg_path is not None:
        attach_svg_fallback(doc, inline_shape, svg_path)


def add_caption(doc, text: str, above=False) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if above else WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = above
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(8 if not above else 3)
    add_inline(p, text, 9.5)


def clean_cell(text: str) -> str:
    return re.sub(r"[`*_]", "", text.strip())


def widths_for_table(headers: list[str]) -> list[int]:
    n = len(headers)
    total = 9360
    if n == 2:
        return [3000, 6360]
    if n == 3:
        return [1700, 3630, 4030]
    if n == 4:
        return [1650, 2450, 2700, 2560]
    if n == 5:
        return [1050, 1760, 2500, 2700, 1350]
    base = total // n
    return [base] * (n - 1) + [total - base * (n - 1)]


def add_markdown_table(doc, lines: list[str]) -> None:
    rows = []
    for line in lines:
        values = [clean_cell(v) for v in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", v.replace(" ", "")) for v in values):
            continue
        rows.append(values)
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for r_idx, values in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, values[c_idx] if c_idx < len(values) else "", 8.5)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(NAVY)
                    set_east_asia(run, "Yu Gothic")
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "FAFBFC")
    repeat_table_header(table.rows[0])
    set_table_geometry(table, widths_for_table(rows[0]))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def create_numbering_instance(doc: Document, ordered: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 0
    num_id = (max(num_ids) + 1) if num_ids else 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "space")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "280")
    ppr.extend([tabs, ind])
    lvl.extend([start, num_fmt, lvl_text, suff, ppr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_paragraph(doc, text: str, numbered: bool, num_id: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    ppr.append(num_pr)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    add_inline(p, text)


def render_markdown(doc: Document, lines: list[str], assets: dict[str, tuple[Path, Path]], main_body=False) -> None:
    i = 0
    current_h2 = ""
    first_h1 = True
    active_num_id = None
    active_list_kind = None
    while i < len(lines):
        raw = lines[i].rstrip()
        text = raw.strip()
        if not text or text == "---":
            i += 1
            continue
        figure = re.fullmatch(r"\[\[FIGURE:(ARCHITECTURE|OPERATION-FLOW|WINDOWS-FIRST-SETUP)\]\]", text)
        if figure:
            key = {
                "ARCHITECTURE": "architecture",
                "OPERATION-FLOW": "operation-flow",
                "WINDOWS-FIRST-SETUP": "windows-first-setup",
            }[figure.group(1)]
            png, svg = assets[key]
            if key == "architecture":
                add_picture_with_alt(doc, png, 6.15, "WPFクライアント、中央サーバー、SMBサーバー、Agent、Gatewayの接続と信頼境界を示す構成図", svg)
                add_caption(doc, "図2 Watashiの構成と主な信頼境界")
            elif key == "operation-flow":
                add_picture_with_alt(doc, png, 6.15, "認証、パス解決、認可、経路選択、SMB実行、監査と応答の順序を示す処理フロー", svg)
                add_caption(doc, "図3 ファイル操作の共通処理フロー")
            else:
                add_picture_with_alt(doc, png, 6.15, "GID入力、Windows認証、本人照合、初回設定待ち状態の確認、本人によるパスワード設定と代替経路を示すフロー", svg)
                add_caption(doc, "図4 初回Windowsユーザー認証とパスワード設定の流れ")
            i += 1
            continue
        screenshot = re.fullmatch(r"\[\[SCREENSHOT:([A-Z-]+)\]\]", text)
        if screenshot:
            key = screenshot.group(1)
            png, svg = assets[key]
            spec = SCREENSHOT_SPECS[key]
            add_picture_with_alt(doc, png, 6.15, spec["alt"], svg, page_break_before=key == "WATASHI-CLIENT")
            add_caption(doc, spec["caption"])
            i += 1
            continue
        if text.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, table_lines)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", text)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 1:
                first_h1 = False
                p = doc.add_heading(title, level=1)
                if main_body and title.startswith("付録"):
                    p.paragraph_format.page_break_before = True
            elif level == 2:
                current_h2 = title
                p = doc.add_heading(title, level=2)
            else:
                p = doc.add_heading(title, level=3)
            for run in p.runs:
                set_east_asia(run, "Yu Gothic")
            i += 1
            active_num_id = None
            active_list_kind = None
            continue
        if re.match(r"^表\d+\s", text):
            add_caption(doc, text, above=True)
            i += 1
            continue
        if text.startswith(">"):
            p = doc.add_paragraph(style="Quote")
            add_inline(p, text.lstrip("> "))
            i += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", text)
        bullet = re.match(r"^[-*]\s+(.+)$", text)
        if numbered:
            if active_list_kind != "ordered":
                active_num_id = create_numbering_instance(doc, True)
            active_list_kind = "ordered"
            add_list_paragraph(doc, numbered.group(1), True, active_num_id)
            i += 1
            continue
        if bullet:
            if active_list_kind != "bullet":
                active_num_id = create_numbering_instance(doc, False)
            active_list_kind = "bullet"
            add_list_paragraph(doc, bullet.group(1), False, active_num_id)
            i += 1
            continue
        active_num_id = None
        active_list_kind = None
        p = doc.add_paragraph()
        add_inline(p, text)
        i += 1


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(46)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TECHNICAL REPORT")
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    set_east_asia(r, "Yu Gothic")

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(22)
    add_inline(p, "既存ツールPANDAの課題を踏まえた\n「Watashi」の設計・実装・検証", 28)
    for run in p.runs:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(NAVY)
        set_east_asia(run, "Yu Gothic")

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "安全性・認可・監査性・運用性向上のアプローチ", 15)
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(TEAL)
        set_east_asia(run, "Yu Gothic")

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    repeat_table_header(table.rows[0])
    set_table_geometry(table, [6900])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    add_inline(p, "PANDA：既存ツール　／　Watashi：後継ツール", 11)
    for run in p.runs:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(88)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "評価基準日　2026年9月9日", 11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "対象　2026年9月9日時点のWatashi技術報告", 10)


def build(source: Path, output: Path) -> None:
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    body_idx = next(i for i, line in enumerate(lines) if re.match(r"^# 1\.", line.strip()))
    abstract_idx = max(i for i, line in enumerate(lines[:body_idx]) if line.strip().startswith("## "))
    abstract_lines = lines[abstract_idx + 1:body_idx]
    body_lines = lines[body_idx:]

    output.parent.mkdir(parents=True, exist_ok=True)
    asset_dir = source.parent / "assets"
    assets = build_diagrams(asset_dir)
    assets.update(build_screenshot_placeholders(asset_dir))

    doc = Document()
    configure_styles(doc)
    configure_page(doc.sections[0])
    doc.sections[0].different_first_page_header_footer = True
    add_running_header(doc.sections[0])
    add_page_number_footer(doc.sections[0])
    doc.core_properties.title = "既存ツールPANDAの課題を踏まえた「Watashi」の設計・実装・検証"
    doc.core_properties.subject = "Watashi 技術レポート"
    doc.core_properties.keywords = "Watashi, PANDA, SMB, CIFS, 認証, 認可, 監査, Agent"
    doc.core_properties.comments = "2026-09-09時点のWatashi技術報告"
    set_update_fields(doc)

    add_cover(doc)
    doc.add_page_break()

    doc.add_heading("概要", level=1)
    render_markdown(doc, abstract_lines, assets)
    doc.add_page_break()

    p = doc.add_heading("目次", level=1)
    p.paragraph_format.space_after = Pt(16)
    add_toc_with_static_result(doc, [
        ("1. 序論", 4),
        ("2. 要求と開発経緯", 6),
        ("3. 設計方針と全体構成", 9),
        ("4. 主要機能の実装", 12),
        ("5. 検証方法", 21),
        ("6. 結果", 22),
        ("7. 考察", 25),
        ("8. 結論と今後の課題", 27),
        ("参考文献", 28),
        ("付録A 要求・実装・検証の対応", 30),
    ])
    doc.add_page_break()

    render_markdown(doc, body_lines, assets, main_body=True)

    for section in doc.sections:
        configure_page(section)
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_report_docx.py SOURCE.md OUTPUT.docx")
    build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
